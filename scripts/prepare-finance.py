"""Read the supplied workbook without modifying it; produce a private import file."""
import argparse
import json
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('workbook')
parser.add_argument('--output', default='outputs/finance-import.json')
args = parser.parse_args()
w = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
transactions = []
for r in w['Master Transaksi'].iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    transactions.append(dict(id=str(r[0]), date=r[1].strftime('%Y-%m-%d'),
        bankPeriod=str(r[2]), period=str(r[3]), bank=str(r[5]), account=str(r[6]),
        description=r[7] or '', reference=str(r[8] or ''), debit=r[9] or 0,
        credit=r[10] or 0, balance=r[11], gl=r[13] or '',
        treatment=r[15] or '', pending=r[16] == 'YA', note=r[17] or '',
        source=r[18] or '', page=r[19]))
coa = [dict(code=r[0], name=r[1], report=r[2], group=r[3], normal=r[4])
       for r in w['COA'].iter_rows(min_row=4, values_only=True) if r[0]]
controls = [dict(period=r[0], bank=r[1], opening=r[2], closing=r[6])
            for r in w['Kontrol Posting'].iter_rows(min_row=5, max_row=10, values_only=True)]
pos_sheet = w['Analisa Pendapatan']
pos = [dict(period=pos_sheet.cell(3,c).value, net=pos_sheet.cell(6,c).value,
            cash=pos_sheet.cell(10,c).value, noncash=pos_sheet.cell(11,c).value)
       for c in range(2,5)]
def snapshot(sheet):
    s = w[sheet]
    return [dict(label=s.cell(r,2).value or s.cell(r,1).value,
                 values={str(s.cell(3,c).value): s.cell(r,c).value for c in range(3,6)})
            for r in range(4, s.max_row+1)
            if any(isinstance(s.cell(r,c).value, (int,float)) for c in range(3,6))]
data = dict(version=1, source=Path(args.workbook).name, transactions=transactions,
            coa=coa, controls=controls, pos=pos,
            balanceSnapshot=snapshot('Neraca Draft'), profitSnapshot=snapshot('Laba Rugi'))
out = Path(args.output)
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(data, ensure_ascii=False, allow_nan=False), encoding='utf-8')
print(f'Prepared {len(transactions)} transactions; {len(controls)} bank controls. Output: {out}')
